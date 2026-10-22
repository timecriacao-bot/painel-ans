"""
Gera o modelo de planilha de comissoes (modelo-comissoes.xlsx) ja preenchido com todas as
operadoras x UF x tipo de contratacao que tem beneficiarios medico-hospitalares no mes mais
recente. Usa so dados publicos da ANS; as comissoes sao preenchidas pela Ramo fora do GitHub.

Chamado pelo ans_update.py (funcao build_modelo). Requer: pip install openpyxl
"""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

F = "Arial"
PURPLE, ORANGE, GREY, YELLOW, INFO = "452780", "E8972C", "EEEEF2", "FFF6CC", "F3F0FA"
TIPOS = ["Individual/familiar", "Coletivo empresarial", "Coletivo por adesão"]
UFS = ["Todas", "AC", "AL", "AM", "AP", "BA", "CE", "DF", "ES", "GO", "MA", "MG", "MS", "MT", "PA", "PB",
       "PE", "PI", "PR", "RJ", "RN", "RO", "RR", "RS", "SC", "SE", "SP", "TO"]
EXTRA_ROWS = 200  # linhas em branco no fim para incluir administradoras/regras extras

# (titulo, largura, tipo, formato, dica) -- tipo: ans = vem da ANS, input = Ramo preenche, calc = formula
COLS = [
    ("Registro ANS", 12, "ans", "@", "Da ANS. Chave para cruzar com o painel."),
    ("Operadora / administradora", 40, "ans", "@", "Da ANS (razão social). Em linhas novas, digite o nome."),
    ("Modalidade", 22, "ans", "@", "Da ANS."),
    ("UF", 7, "ans", "@", "Da ANS. Em linhas novas, escolha na lista (ou “Todas”)."),
    ("Tipo de plano", 22, "ans", "@", "Da ANS. Em linhas novas, escolha na lista."),
    ("Beneficiários (ANS)", 13, "ans", "#,##0", "Beneficiários médico-hospitalares ativos nessa UF e tipo, no mês indicado na aba Instruções."),
    ("Comissão total (% de 1 mensalidade)", 16, "input", "0%", "Soma de todas as parcelas. Ex.: 3 parcelas de 100% = 300%. Deixe em branco se não vende."),
    ("Parcelas (meses)", 10, "input", "0", "Opcional. Em quantos meses a comissão é paga."),
    ("Comissão vitalícia (% ao mês)", 12, "input", "0.0%", "Opcional. Não entra no cálculo de retorno."),
    ("Ticket médio por vida (R$/mês)", 14, "input", '"R$" #,##0.00', "Mensalidade média por beneficiário vendida."),
    ("Vidas por contrato", 10, "input", "0.0", "Média de vidas por venda."),
    ("CPL médio (R$)", 11, "input", '"R$" #,##0.00', "Custo por lead das campanhas."),
    ("Conversão lead → venda (%)", 12, "input", "0.0%", "% dos leads que viraram contrato."),
    ("Vigência da tabela (mm/aaaa)", 12, "input", "@", "Ex.: 10/2026."),
    ("Fonte / observação", 28, "input", "@", "Tabela da operadora, contrato, histórico de campanhas..."),
    ("Comissão por venda (R$)", 14, "calc", '"R$" #,##0.00;-"R$" #,##0.00;-', "Calculado: Comissão total × Ticket × Vidas."),
    ("Receita por lead (R$)", 13, "calc", '"R$" #,##0.00;-"R$" #,##0.00;-', "Calculado: Comissão por venda × Conversão."),
    ("Retorno por R$ 1 de mídia", 12, "calc", '0.00"x";-0.00"x";-', "Calculado: Receita por lead ÷ CPL. Acima de 1,00x a campanha se paga com a comissão inicial."),
]


def _instrucoes(ws, competencia, n):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 112
    lines = [
        ("t", "Modelo de comissões por operadora · Ramo Hub"),
        ("s", f"Lista gerada automaticamente com dados abertos da ANS · competência {competencia} · {n} linhas (operadora × UF × tipo de plano)."),
        ("", ""),
        ("h", "O que já vem preenchido (colunas lilás, dados da ANS)"),
        ("p", "Registro ANS, operadora, modalidade, UF, tipo de plano e beneficiários médico-hospitalares no mês. Só entram combinações com beneficiários ativos."),
        ("p", "Ter beneficiários num estado não garante que a operadora venda ali hoje: confira com a tabela comercial."),
        ("", ""),
        ("h", "Como preencher (colunas amarelas)"),
        ("p", "1. Preencha só as linhas que vocês vendem ou anunciam. Linhas com a comissão em branco são ignoradas pelo painel."),
        ("p", "2. Comissão total = soma das parcelas em % de UMA mensalidade. Ex.: 3 parcelas de 100% = 300%."),
        ("p", "3. Sem ticket, vidas, CPL e conversão o retorno não é calculado. Pode começar só com a comissão."),
        ("p", "4. Use o filtro do cabeçalho para achar UF, tipo ou operadora. Ordem padrão: UF e, dentro dela, mais beneficiários primeiro."),
        ("p", "5. Administradoras de benefícios e regras que valem para o Brasil todo: use as linhas em branco no fim, com UF “Todas”."),
        ("p", "6. A linha 4 é um EXEMPLO com valores fictícios. Apague antes de enviar."),
        ("", ""),
        ("h", "Cálculos automáticos (colunas cinza)"),
        ("p", "• Comissão por venda (R$) = Comissão total × Ticket médio × Vidas por contrato."),
        ("p", "• Receita por lead (R$) = Comissão por venda × Conversão."),
        ("p", "• Retorno por R$ 1 de mídia = Receita por lead ÷ CPL. Acima de 1,00x a campanha se paga só com a comissão inicial (sem contar a vitalícia)."),
        ("", ""),
        ("h", "Privacidade"),
        ("p", "Este arquivo em branco é público (vem do GitHub). DEPOIS de preenchido com comissões, NÃO publique no GitHub: "
              "ele vai só para o plugin do WordPress, visível apenas para usuários logados."),
    ]
    for r, (kind, txt) in enumerate(lines, 2):
        c = ws.cell(row=r, column=2, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = {"t": Font(name=F, size=16, bold=True, color=PURPLE),
                  "s": Font(name=F, size=10, italic=True, color="6B6480"),
                  "h": Font(name=F, size=12, bold=True, color=ORANGE)}.get(kind, Font(name=F, size=10))


def build_modelo(rows, competencia, out_path):
    """rows: lista de dicts com reg, nome, mod, uf, tipo, ben (ja ordenada)."""
    thin = Side(style="thin", color="D5D0E3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    _instrucoes(wb.active, competencia, len(rows))
    wb.active.title = "Instruções"

    ls = wb.create_sheet("Listas")
    ls["A1"], ls["B1"] = "Tipo de plano", "UF"
    for i, t in enumerate(TIPOS + ["Todos"], 2):
        ls.cell(row=i, column=1, value=t)
    for i, u in enumerate(UFS, 2):
        ls.cell(row=i, column=2, value=u)
    ls.sheet_state = "hidden"

    ws = wb.create_sheet("Comissões", 1)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Comissões por operadora · preencha as células amarelas das linhas que vocês vendem (linha 4 é exemplo fictício)"
    ws["A1"].font = Font(name=F, size=12, bold=True, color=PURPLE)
    ws["A2"] = f"Legenda:  lilás = dados da ANS ({competencia})  ·  amarelo = você preenche  ·  cinza = calcula sozinho"
    ws["A2"].font = Font(name=F, size=9, italic=True, color="6B6480")
    HR, FIRST = 3, 4
    fills = {"ans": INFO, "input": YELLOW, "calc": GREY}
    heads = {"ans": "6E5A9E", "input": PURPLE, "calc": "6B6480"}
    for j, (name, w, kind, fmt, tip) in enumerate(COLS, 1):
        c = ws.cell(row=HR, column=j, value=name)
        c.font = Font(name=F, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=heads[kind])
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = border
        c.comment = Comment(tip, "Ramo Hub")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[HR].height = 50
    ws.freeze_panes = "C4"

    example = {"reg": "", "nome": "EXEMPLO FICTÍCIO (apague esta linha)", "mod": "", "uf": "SP",
               "tipo": "Coletivo empresarial", "ben": None,
               "inp": [3.0, 3, 0.02, 550, 3, 45, 0.08, "10/2026", "Valores ilustrativos, não são dados reais"]}
    data = [example] + list(rows)
    last = FIRST + len(data) - 1 + EXTRA_ROWS
    styles = [(Font(name=F, size=10, color="0000FF" if k == "input" else "000000"),
               PatternFill("solid", fgColor=fills[k]), fmt) for (_, _, k, fmt, _) in COLS]
    for i in range(FIRST, last + 1):
        for j, (font, fill, fmt) in enumerate(styles, 1):
            c = ws.cell(row=i, column=j)
            c.font, c.fill, c.number_format, c.border = font, fill, fmt, border
        ws.cell(row=i, column=16, value=f'=IF(OR(G{i}="",J{i}="",K{i}=""),"",G{i}*J{i}*K{i})')
        ws.cell(row=i, column=17, value=f'=IF(OR(P{i}="",M{i}=""),"",P{i}*M{i})')
        ws.cell(row=i, column=18, value=f'=IF(OR(Q{i}="",L{i}="",L{i}=0),"",Q{i}/L{i})')
    for i, r in enumerate(data, FIRST):
        for j, v in enumerate([r["reg"], r["nome"], r["mod"], r["uf"], r["tipo"], r["ben"]], 1):
            ws.cell(row=i, column=j, value=v)
        for j, v in enumerate(r.get("inp", []), 7):
            ws.cell(row=i, column=j, value=v)
    for j in range(1, 16):
        ws.cell(row=FIRST, column=j).font = Font(name=F, size=10, italic=True, color="0000FF")

    dvs = [
        (DataValidation(type="list", formula1=f"=Listas!$A$2:$A${len(TIPOS) + 2}", allow_blank=True,
                        errorTitle="Tipo de plano", error="Escolha uma opção da lista."), f"E{FIRST}:E{last}"),
        (DataValidation(type="list", formula1=f"=Listas!$B$2:$B${len(UFS) + 1}", allow_blank=True,
                        errorTitle="UF", error="Escolha uma UF da lista ou “Todas”."), f"D{FIRST}:D{last}"),
        (DataValidation(type="decimal", operator="between", formula1="0", formula2="20", allow_blank=True,
                        errorTitle="Comissão total", error="Em % de uma mensalidade (ex.: 300%)."), f"G{FIRST}:G{last}"),
        (DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                        errorTitle="Percentual", error="Digite um percentual entre 0% e 100%."), f"I{FIRST}:I{last} M{FIRST}:M{last}"),
    ]
    for dv, rng in dvs:
        ws.add_data_validation(dv)
        for part in rng.split():
            dv.add(part)
    ws.auto_filter.ref = f"A{HR}:R{last}"
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(out_path)
    return len(rows)
